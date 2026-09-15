from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path

OUT = Path(__file__).parent / 'data' / 'store.xlsx'
wb = Workbook()
black='111111'; gold='C9A452'; cream='F8F3EA'; tan='E8DCCB'; gray='6B655D'; white='FFFFFF'
thin=Side(style='thin', color='D7C8B5')

def style_sheet(ws, widths):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 36
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor=black)
        cell.font = Font(color=gold, bold=True, size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=Side(style='medium', color=gold))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
            cell.border = Border(bottom=thin)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='FCF9F4')
    for i,w in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.sheet_properties.tabColor = gold

products = wb.active
products.title='المنتجات'
headers=['معرّف المنتج','اسم المنتج','الفئة','السعر','السعر قبل الخصم','وصف قصير','الوصف الكامل','الصورة الرئيسية','صور إضافية','متوفر','مميز','وصل حديثًا','الترتيب','الخامة','المقاس']
products.append(headers)
rows=[
['bidaya-black','دفتر بداية — أسود','دفاتر',295,340,'دفتر أنيق لمساحتك الخاصة وأفكارك اليومية.','غلاف صلب مطفي، ورق كريمي مريح للعين، فاصل قماشي، وتصميم عملي يناسب الدراسة والعمل والكتابة اليومية.','notebook-bidaya.jpg','', 'نعم','نعم','نعم',1,'غلاف صلب وورق كريمي','A5'],
['khotwa-weekly','بلانر خطوة الأسبوعي','تنظيم وتخطيط',220,'','أسبوعك كامل قدامك بخطوات واضحة.','بلانر أسبوعي بتقسيم بسيط يساعدك ترتب أولوياتك ومواعيدك وملاحظاتك من غير زحمة.','planner-khotwa.jpg','', 'نعم','نعم','نعم',2,'ورق سميك 120 جم','A4'],
['athar-gratitude','مفكرة أثر للامتنان','دفاتر',260,'','مساحة هادئة لتوثيق الأشياء الجميلة.','مفكرة يومية خفيفة بأسئلة بسيطة للامتنان والتأمل، بتصميم دافئ وتجليد متين.','gratitude-athar.jpg','', 'نعم','نعم','لا',3,'غلاف قماشي وورق كريمي','A5'],
['satr-pen','قلم سطر المعدني','أقلام',135,'','قلم متوازن بلمسة سوداء وذهبية.','قلم معدني أنيق بسن ناعم ومشبك ذهبي، مناسب للاستخدام اليومي والهدايا.','pen-satr.jpg','', 'نعم','لا','نعم',4,'معدن مطفي','0.7 مم'],
['lahza-box','بوكس لحظة','بوكسات وهدايا',595,650,'هدية متكاملة ومغلفة بعناية.','بوكس هدايا يضم مفكرة، قلمًا، كروت تخطيط وفواصل مختارة، مع تغليف أنيق جاهز للإهداء.','gift-lahza.jpg','', 'نعم','نعم','نعم',5,'تشكيلة متنوعة','بوكس متوسط'],
['bayn-bookmarks','فواصل بين السطور','فواصل وستيكرز',95,'','ثلاثة فواصل تضيف لمستك لكل كتاب.','مجموعة فواصل بطباعة عالية الجودة وألوان محايدة، مع خيط أنيق وتشطيب مقاوم للاستخدام.','bookmarks-bayn.jpg','', 'نعم','لا','لا',6,'ورق مقوى مغلف','3 قطع'],
['small-details-stickers','ستيكرز تفاصيل صغيرة','فواصل وستيكرز',85,'','ورقة ستيكرز محايدة للتزيين والتخطيط.','ملصقات ورقية بأشكال نباتية وتجريدية تناسب الدفاتر والبلانرز والهدايا.','bookmarks-bayn.jpg','', 'نعم','لا','نعم',7,'ورق لاصق مطفي','ورقة A5'],
['fekra-note-pad','نوت باد فكرة','تنظيم وتخطيط',120,'','كل أفكارك السريعة في مكان واحد.','نوت باد يومي للمهام والأفكار والأولويات، سهل الحمل والاستخدام على المكتب.','planner-khotwa.jpg','', 'نعم','لا','لا',8,'50 ورقة 100 جم','A5'],
]
for r in rows: products.append(r)
style_sheet(products,[19,28,20,12,19,35,55,27,34,12,12,15,10,24,18])
products.auto_filter.ref=f'A1:O{products.max_row}'
yes_no=DataValidation(type='list', formula1='"نعم,لا"', allow_blank=False)
products.add_data_validation(yes_no); yes_no.add('J2:L1000')

categories=wb.create_sheet('الفئات')
categories.append(['معرّف الفئة','اسم الفئة','صورة الفئة','وصف قصير','الترتيب'])
for r in [
['new','وصل حديثًا','gift-lahza.jpg','أحدث القطع والمجموعات',1],
['notebooks','دفاتر','notebook-bidaya.jpg','مساحات للأفكار والحكايات',2],
['planning','تنظيم وتخطيط','planner-khotwa.jpg','أدوات ليوم أوضح',3],
['pens','أقلام','pen-satr.jpg','تفاصيل تكمّل مكتبك',4],
['gifts','بوكسات وهدايا','gift-lahza.jpg','هدايا جاهزة تصنع ذكرى',5],
['accessories','فواصل وستيكرز','bookmarks-bayn.jpg','لمسات صغيرة بطابعك',6],
]: categories.append(r)
style_sheet(categories,[20,25,28,38,12])

settings=wb.create_sheet('إعدادات المتجر')
settings.append(['الإعداد','القيمة'])
for r in [
['اسم المتجر','أثر'],
['رقم واتساب','201030263241'],
['العملة','ج.م'],
['شريط الإعلان','شحن لكل محافظات مصر  •  تأكيد الطلب بسهولة عبر واتساب'],
['عنوان الواجهة','تفاصيل صغيرة، تصنع أثرًا كبيرًا'],
['وصف الواجهة','دفاتر، أدوات تخطيط وهدايا مصممة لترافق أيامك وتترك لمسة لا تُنسى.'],
['صورة الواجهة','hero-athar.jpg'],
['نص زر الواجهة','اكتشف المجموعة'],
['عنوان قسم الفئات','اختار اللي يناسب يومك'],
['عنوان المنتجات المميزة','اختيارات صنعت لتبقى'],
['عنوان وصل حديثًا','وصل حديثًا'],
['عنوان آراء العملاء','كلام ترك أثرًا'],
['عنوان من نحن','كل تفصيلة لها معنى'],
['نص من نحن','في أثر نختار كل قطعة بهدوء واهتمام؛ من ملمس الورق حتى آخر لمسة في التغليف، علشان توصلك حاجة بسيطة لكن تفضل معاك.'],
['إنستجرام',''],
['فيسبوك',''],
['بريد إلكتروني',''],
['ملاحظة الشحن','تكلفة الشحن تُحسب حسب المحافظة، ويتم تأكيد الطلب والموعد عبر واتساب.'],
['رسالة أسفل الصفحة','مصنوع بحب ليترك أثرًا جميلًا.'],
]: settings.append(r)
style_sheet(settings,[34,90])
settings.auto_filter.ref='A1:B1'

shipping=wb.create_sheet('الشحن')
shipping.append(['المحافظة','التكلفة','متاح'])
shipping_rows=[
('القاهرة',60),('الجيزة',65),('القليوبية',65),('الإسكندرية',75),('البحيرة',75),('الغربية',75),('الدقهلية',75),('الشرقية',75),('المنوفية',75),('كفر الشيخ',80),('دمياط',80),('بورسعيد',85),('الإسماعيلية',85),('السويس',85),('الفيوم',80),('بني سويف',80),('المنيا',90),('أسيوط',90),('سوهاج',95),('قنا',100),('الأقصر',100),('أسوان',110),('البحر الأحمر',110),('الوادي الجديد',120),('مطروح',110),('شمال سيناء',120),('جنوب سيناء',120)
]
for gov,cost in shipping_rows: shipping.append([gov,cost,'نعم'])
style_sheet(shipping,[25,15,12])
available=DataValidation(type='list', formula1='"نعم,لا"', allow_blank=False)
shipping.add_data_validation(available); available.add('C2:C100')

reviews=wb.create_sheet('آراء العملاء')
reviews.append(['الاسم','الرأي','التقييم','الترتيب','ظاهر'])
for r in [
['سارة','التغليف كان أنيق جدًا وكل تفصيلة في الأوردر واضحة إنها معمولة باهتمام.',5,1,'نعم'],
['مريم','الخامة أجمل من الصور والطلب وصل مرتب. أكيد مش آخر مرة.',5,2,'نعم'],
['نور','البلانر بسيط وعملي وخلاني أرتب أسبوعي من غير ما أحس بزحمة.',5,3,'نعم'],
]: reviews.append(r)
style_sheet(reviews,[20,72,12,12,12])
show=DataValidation(type='list', formula1='"نعم,لا"', allow_blank=False)
reviews.add_data_validation(show); show.add('E2:E1000')

instructions=wb.create_sheet('اقرأني أولًا')
instructions.sheet_view.rightToLeft=True
instructions.sheet_properties.tabColor=gold
instructions.column_dimensions['A'].width=120
instructions['A1']='إدارة متجر أثر من ملف Excel'
instructions['A1'].font=Font(name='Arial',size=20,bold=True,color=gold)
instructions['A1'].fill=PatternFill('solid',fgColor=black)
instructions['A1'].alignment=Alignment(horizontal='center',vertical='center')
instructions.row_dimensions[1].height=44
notes=[
'1) عدّل المنتجات والأسعار والفئات والنصوص من الشيتات، ولا تغيّر أسماء الأعمدة أو أسماء الشيتات.',
'2) لإضافة صورة: ارفعها من لوحة الإدارة، ثم اكتب اسم الملف بالامتداد في خانة الصورة الرئيسية، مثال: product-1.jpg.',
'3) أسماء الصور الإضافية تُكتب مفصولة بفاصلة: image-2.jpg, image-3.jpg.',
'4) القيم نعم/لا تتحكم في ظهور المنتج وحالات مميز ووصل حديثًا.',
'5) بعد الانتهاء احفظ الملف بصيغة XLSX، ثم ارفعه من لوحة الإدارة. الموقع يقرأ التغييرات فورًا.',
'6) أسعار الشحن في الملف تجريبية؛ راجعها وعدّلها قبل إطلاق المتجر.',
'7) الآراء الحالية أمثلة تجريبية؛ استبدلها بآراء حقيقية قبل النشر.',
'8) معرّف المنتج يجب أن يكون فريدًا، ويفضل بالإنجليزية ومن دون مسافات مثل: black-notebook.',
]
for i,n in enumerate(notes,3):
    c=instructions.cell(i,1,n); c.alignment=Alignment(horizontal='right',vertical='top',wrap_text=True); c.font=Font(size=12,color=black)
    instructions.row_dimensions[i].height=30
instructions['A13']='ترتيب الرفع المقترح: الصور أولًا ← ملف Excel بعد التعديل ← افتح المتجر وراجع النتيجة.'
instructions['A13'].font=Font(bold=True,color='8B6A25',size=12)
instructions['A13'].alignment=Alignment(horizontal='right',wrap_text=True)

wb.active=0
OUT.parent.mkdir(parents=True,exist_ok=True)
wb.save(OUT)
print(OUT)
